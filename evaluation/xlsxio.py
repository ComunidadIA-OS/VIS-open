"""Hojas de anotación en Excel (.xlsx) y vuelta a CSV.

Los CSV de `data/anotacion/` van en UTF-8 y separados por comas. Excel en español los abre mal con
doble clic: todo queda en una columna y los acentos se rompen. Para anotar es más cómodo un .xlsx
con las columnas separadas, el texto ajustado, la cabecera fija y listas desplegables en las
columnas que se rellenan (`es_pregunta`: si/no; `label`: mujer/hombre/desconocido).

    python -m evaluation.xlsxio a-xlsx data/anotacion/preguntas_para_anotar.csv   # -> .xlsx
    python -m evaluation.xlsxio a-csv  data/anotacion/preguntas_para_anotar.xlsx  # -> .csv

Cuando la persona devuelve el .xlsx relleno, `a-csv` lo pasa a CSV UTF-8 con comas y el resto de
herramientas (`evaluation.questions_eval`, `evaluation.benchmark`...) lo leen como siempre.
Requiere `openpyxl` (extra `[anotacion]`, MIT).
"""

from __future__ import annotations

import argparse
import csv
from collections.abc import Sequence
from pathlib import Path

DROPDOWNS = {
    "es_pregunta": ["si", "no"],
    "label": ["mujer", "hombre", "desconocido"],
}
WIDTHS = {
    "turn_id": 9, "t_inicio": 10, "context_before": 60, "text": 90, "frase": 50,
    "cargo_destinatario": 34, "label": 14, "es_pregunta": 14, "video_id": 10, "event_id": 14,
    "genero_sistema": 16,
}  # fmt: skip
WRAPPED = {"context_before", "text", "frase", "cargo_destinatario"}


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - mensaje claro para quien no lo tiene
        raise SystemExit("Falta openpyxl: pip install -e '.[anotacion]'") from exc
    return openpyxl


def csv_to_xlsx(csv_path: Path, xlsx_path: Path | None = None) -> Path:
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    out = xlsx_path or csv_path.with_suffix(".xlsx")
    with csv_path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        raise ValueError(f"{csv_path} está vacío")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "hoja"
    for row in rows:
        ws.append(row)
    header = rows[0]
    fill = PatternFill("solid", fgColor="1F6F7D")
    for col, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS.get(name, 18)
        for r in range(2, len(rows) + 1):
            c = ws.cell(row=r, column=col)
            c.alignment = Alignment(wrap_text=name in WRAPPED, vertical="top")
        if name in DROPDOWNS:
            dv = DataValidation(type="list", formula1='"' + ",".join(DROPDOWNS[name]) + '"',
                                allow_blank=True)  # fmt: skip
            dv.error, dv.errorTitle = "Elige un valor de la lista", "Valor no válido"
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{len(rows)}")
            for r in range(2, len(rows) + 1):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FFF4CC")
    ws.freeze_panes = "A2"
    wb.save(out)
    return out


def xlsx_to_csv(xlsx_path: Path, csv_path: Path | None = None) -> Path:
    openpyxl = _openpyxl()
    out = csv_path or xlsx_path.with_suffix(".csv")
    ws = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True).active
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            if not any(v not in (None, "") for v in row):
                continue
            w.writerow(["" if v is None else str(v).strip() for v in row])
    return out


def main(argv: Sequence[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("a-xlsx", "a-csv"):
        s = sub.add_parser(name)
        s.add_argument("files", type=Path, nargs="+")
    args = ap.parse_args(argv)
    for f in args.files:
        out = csv_to_xlsx(f) if args.cmd == "a-xlsx" else xlsx_to_csv(f)
        print(f"{f} -> {out}")


if __name__ == "__main__":
    main()
