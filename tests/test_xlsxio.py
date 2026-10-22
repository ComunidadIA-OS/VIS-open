import csv

import pytest

pytest.importorskip("openpyxl")
from openpyxl import load_workbook  # noqa: E402

from evaluation.xlsxio import csv_to_xlsx, xlsx_to_csv  # noqa: E402

ROWS = [
    ["turn_id", "context_before", "text", "es_pregunta", "frase"],
    ["Q001", "Gracias, señora ministra.", "¿Qué piensa hacer, señor presidente?", "", ""],
    ["Q002", "", 'Dijo: "no, nunca", y se fue; ¡ñandú!', "", ""],
]


def _write(path):
    with path.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(ROWS)


def test_csv_to_xlsx_has_separate_columns_and_dropdowns(tmp_path):
    src = tmp_path / "hoja.csv"
    _write(src)
    wb = load_workbook(csv_to_xlsx(src))
    ws = wb.active
    assert [c.value for c in ws[1]] == ROWS[0]
    assert ws["C2"].value == "¿Qué piensa hacer, señor presidente?"
    assert ws.freeze_panes == "A2"
    assert any("si,no" in dv.formula1 for dv in ws.data_validations.dataValidation)


def test_roundtrip_through_excel_keeps_the_text(tmp_path):
    src = tmp_path / "hoja.csv"
    _write(src)
    back = xlsx_to_csv(csv_to_xlsx(src), tmp_path / "vuelta.csv")
    with back.open(newline="", encoding="utf-8") as fh:
        assert list(csv.reader(fh)) == ROWS


def test_label_sheet_has_instructions_and_roundtrip_ignores_them(tmp_path):
    src = tmp_path / "t.csv"
    src.write_text("turn_id,label\nT001,\n", encoding="utf-8")
    xlsx = csv_to_xlsx(src)
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["hoja", "Instrucciones"]
    assert any("desconocido" in str(c.value) for c in wb["Instrucciones"]["A"])
    assert xlsx_to_csv(xlsx).read_text(encoding="utf-8").splitlines() == ["turn_id,label", "T001,"]
